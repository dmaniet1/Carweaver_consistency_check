import json
from collections import OrderedDict, defaultdict
from openpyxl import load_workbook

INPUT_FILE = "your_excel_file.xlsx" #Your excel file
OUTPUT_FILE = "VE_qa1_test_cases_flat_.json"
CONFLICT_LOG = "tc_conflicts.log"

wb = load_workbook(INPUT_FILE, read_only=True, data_only=True)
ws = wb["Test Cases"]

# Header row is row 6
headers = [cell.value for cell in next(ws.iter_rows(min_row=6, max_row=6))]
idx = {header: i for i, header in enumerate(headers) if header}


def normalize_automation(value):
    """
    Normalize Test Automation values.
    """

    if value is None:
        return "Not automated"

    value = str(value).strip()

    if value == "":
        return "Not automated"

    if value.lower() == "not set":
        return "Not automated"

    if value.lower() == "yes":
        return "Automated"

    if value.lower() == "no":
        return "Not automated"

    return value


test_cases = OrderedDict()
duplicate_check = defaultdict(set)

for row in ws.iter_rows(min_row=7, values_only=True):

    if all(v is None for v in row):
        continue

    tc_xhandle = row[idx["TC x-Handle"]]

    if tc_xhandle is None:
        continue

    tc_xhandle = str(tc_xhandle)

    name = row[idx["Test Case"]]

    automation = normalize_automation(
        row[idx["Test Automation"]]
    )

    duplicate_check[tc_xhandle].add(
        (name, automation)
    )

    test_cases.setdefault(
        tc_xhandle,
        {
            "xHandle": tc_xhandle,
            "Name": name,
            "Test Automation Attribute": automation
        }
    )

#
# Conflicts
#
conflicting_duplicates = {
    k: list(v)
    for k, v in duplicate_check.items()
    if len(v) > 1
}

#
# Print conflicts 
#
if conflicting_duplicates:

    print(
        f"\nWARNING: Found "
        f"{len(conflicting_duplicates)} "
        f"TC xHandles with conflicting values.\n"
    )

    with open(CONFLICT_LOG, "w", encoding="utf-8") as log:

        for handle, values in conflicting_duplicates.items():

            print(f"TC xHandle: {handle}")

            log.write(f"\nTC xHandle: {handle}\n")

            for idx_value, value in enumerate(values, start=1):

                print(f"  Option {idx_value}:")
                print(f"    Name       : {value[0]}")
                print(f"    Automation : {value[1]}")

                log.write(
                    f"Option {idx_value}: "
                    f"Name='{value[0]}', "
                    f"Automation='{value[1]}'\n"
                )

            print("  -> Fallback: using first occurrence\n")

            log.write(
                "Fallback: using first occurrence\n"
            )

    print(
        f"\nConflict details written to "
        f"{CONFLICT_LOG}\n"
    )

#
# Export JSON
#
with open(
    OUTPUT_FILE,
    "w",
    encoding="utf-8"
) as f:

    json.dump(
        list(test_cases.values()),
        f,
        ensure_ascii=False,
        indent=2
    )

wb.close()

print("\nExport completed successfully")
print(f"Unique TC objects: {len(test_cases)}")
print(f"JSON file: {OUTPUT_FILE}")